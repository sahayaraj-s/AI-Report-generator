"""
schema_detector.py - Generic Schema Detection Layer
Runs FIRST on every uploaded Excel file, before loader.py.
"""
from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import unicodedata
from difflib import SequenceMatcher
from typing import Any

import openpyxl

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Normalisation helpers
# ---------------------------------------------------------------------------

def _norm(s: Any) -> str:
    """Lowercase, strip, collapse whitespace, remove punctuation."""
    if s is None:
        return ""
    t = str(s).strip()
    t = unicodedata.normalize("NFKD", t)
    t = t.lower()
    t = re.sub(r"[^\w\s]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def _sim(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()


def _fuzzy_match_any(text: str, keywords: list, threshold: float = 0.82) -> bool:
    n = _norm(text)
    for kw in keywords:
        nk = _norm(kw)
        if nk in n or n in nk:
            return True
        if _sim(n, nk) >= threshold:
            return True
    return False


# ---------------------------------------------------------------------------
# Column-category keyword sets
# ---------------------------------------------------------------------------

NAME_KEYWORDS = ["name", "student name", "full name", "candidate name", "studentname",
                 "name of student", "name of candidate"]

ID_KEYWORDS = ["enrolment no", "enrolment number", "enrollment no", "enrollment number",
               "roll no", "roll number", "student id", "reg no", "registration no",
               "sno", "sr no", "serial no", "sl no", "admission no"]

PROFILE_KEYWORDS = ["gender", "age", "dob", "date of birth", "qualification", "college",
                    "address", "district", "email", "father",
                    "mother", "occupation", "reference", "personal no", "parents no",
                    "present address", "permanent address", "college name"]

ASSESSMENT_KEYWORDS = ["out of", "max marks", "total marks", "marks obtained",
                       "theory", "practical", "ms word", "ms excel", "basic computer",
                       "basic english", "business writing", "hospital administration",
                       "role play", "soft skill"]

SKILL_KEYWORDS = ["communication", "leadership", "creativity", "technical skill",
                  "analytical", "typing", "wpm", "out of 10",
                  "innovation", "problem solving"]

PLACEMENT_KEYWORDS = ["designation", "organization", "company", "salary", "package",
                      "placed", "placement", "employer", "ctc", "stipend"]

OTHER_SKILLS_KEYWORDS = ["sports", "hobbies", "other skills", "extra activities",
                         "extracurricular"]

IRRELEVANT_KEYWORDS = ["uniform", "shirt", "trouser", "tie", "shoe size", "socks",
                       "chudidhar", "shawl", "dress"]


# ---------------------------------------------------------------------------
# Header row detection
# ---------------------------------------------------------------------------

def _is_blank_or_merged_label(row_vals: list, threshold: int = 3) -> bool:
    non_empty = [v for v in row_vals if v is not None and str(v).strip() != ""]
    return len(non_empty) < threshold


def _is_numeric(v: Any) -> bool:
    try:
        float(str(v).replace(",", ""))
        return True
    except (ValueError, TypeError):
        return False


def _looks_like_header(row_vals: list) -> bool:
    text_cells = [v for v in row_vals if v is not None and not _is_numeric(v)]
    if len(text_cells) < 2:
        return False
    distinct = len(set(str(v).strip().lower() for v in text_cells if str(v).strip()))
    return distinct >= 2


def _detect_header_row(ws, max_scan: int = 20) -> int:
    rows = []
    for row in ws.iter_rows(max_row=max_scan, values_only=True):
        rows.append(list(row))

    for i, row in enumerate(rows):
        if _is_blank_or_merged_label(row):
            continue
        if _looks_like_header(row):
            return i
    return 0


def _detect_second_header_row(ws, header_row_idx: int, max_scan: int = 20):
    rows = []
    for row in ws.iter_rows(max_row=max_scan, values_only=True):
        rows.append(list(row))

    candidate_idx = header_row_idx + 1
    if candidate_idx >= len(rows):
        return None

    row = rows[candidate_idx]
    out_of_count = sum(
        1 for v in row
        if v is not None and re.search(r"out\s+of\s*\d+", str(v), re.IGNORECASE)
    )
    none_count = sum(1 for v in row if v is None or str(v).strip() == "")

    if out_of_count >= 1 or none_count > (len(row) * 0.5):
        return candidate_idx
    return None


# ---------------------------------------------------------------------------
# Column identification
# ---------------------------------------------------------------------------

def _find_name_column(headers: list) -> str | None:
    for h in headers:
        if _fuzzy_match_any(h, NAME_KEYWORDS):
            return h
    return None


def _find_id_column(headers: list) -> str | None:
    for h in headers:
        if _fuzzy_match_any(h, ID_KEYWORDS):
            return h
    return None


# ---------------------------------------------------------------------------
# Sheet category classifier
# ---------------------------------------------------------------------------

def _count_date_like(headers: list) -> int:
    count = 0
    for h in headers:
        if re.search(r"\d{1,2}[/\-\.]\d{1,2}([/\-\.]\d{2,4})?", str(h)):
            count += 1
        elif re.search(r"^\d{1,2}$", str(h).strip()):
            count += 1
    return count


def _classify_sheet(headers: list, sample_values: list, sheet_name: str) -> tuple:
    nh = [_norm(h) for h in headers]
    sheet_norm = _norm(sheet_name)

    # Hostelite / parent meet (check BEFORE irrelevant)
    if any(kw in sheet_norm for kw in ["hostelite", "hostel"]):
        return "hostelite", 0.85
    if any(kw in sheet_norm for kw in ["parent", "parents meet", "parents meet"]):
        return "parent_meet", 0.85

    # Skill matrix — check BEFORE irrelevant because keywords overlap
    skill_hits_early = sum(1 for n in nh if any(kw in n for kw in SKILL_KEYWORDS))
    out_of_10_early = sum(1 for n in nh if re.search(r"out\s+of\s*10", n))
    if out_of_10_early >= 2 or (skill_hits_early >= 3 and out_of_10_early >= 1):
        return "skill_matrix", 0.90

    # Irrelevant (uniform, dress sizes, etc.)
    irr_score = sum(1 for kw in IRRELEVANT_KEYWORDS if any(kw in n for n in nh))
    if irr_score >= 2 or any(kw in sheet_norm for kw in ["uniform"]):
        return "irrelevant", 0.90

    # Attendance: many date-like headers
    date_count = _count_date_like(headers)
    if date_count > 5:
        return "attendance", 0.90
    att_kw_hits = sum(1 for n in nh if any(kw in n for kw in ["attendance", "present", "absent"]))
    if att_kw_hits >= 1 and date_count >= 2:
        return "attendance", 0.85

    # Detect P/A values pattern in data
    flat_vals = [str(v).strip().upper() for row in sample_values for v in row if v is not None]
    pa_ratio = sum(1 for v in flat_vals if v in ("P", "A", "-")) / max(len(flat_vals), 1)
    if pa_ratio > 0.3:
        return "attendance", 0.88

    # Placement
    place_hits = sum(1 for n in nh if any(kw in n for kw in ["designation", "organization", "salary",
                                                               "company", "package", "placed", "employer"]))
    if place_hits >= 2:
        return "placement", 0.88

    # Assessment
    assessment_hits = sum(1 for n in nh if any(kw in n for kw in ASSESSMENT_KEYWORDS))
    if assessment_hits >= 2:
        return "assessment", 0.85

    # Skill matrix (secondary check)
    skill_hits = sum(1 for n in nh if any(kw in n for kw in SKILL_KEYWORDS))
    if skill_hits >= 2:
        out_of_10 = sum(1 for n in nh if re.search(r"out\s+of\s*10", n))
        if out_of_10 >= 1 or skill_hits >= 3:
            return "skill_matrix", 0.85
        return "assessment", 0.75

    # Other skills
    other_hits = sum(1 for n in nh if any(kw in n for kw in OTHER_SKILLS_KEYWORDS))
    if other_hits >= 1:
        return "other_skills", 0.80

    # Profile
    profile_hits = sum(1 for n in nh if any(kw in n for kw in PROFILE_KEYWORDS))
    if profile_hits >= 3:
        return "profile", 0.85
    if profile_hits >= 1:
        return "profile", 0.65

    return "unknown", 0.30


# ---------------------------------------------------------------------------
# Score-column extraction
# ---------------------------------------------------------------------------

def _extract_max_from_text(text: str):
    m = re.search(r"(?:out\s*of|/|\()\s*(\d+(?:\.\d+)?)", str(text), re.IGNORECASE)
    if m:
        return float(m.group(1))
    return None


def _clean_display_name(raw: str) -> str:
    name = re.sub(r"\(?out\s+of\s*\d+(?:\.\d+)?\)?", "", raw, flags=re.IGNORECASE)
    name = re.sub(r"/\s*\d+(?:\.\d+)?", "", name)
    name = re.sub(r"\(\s*\d+(?:\.\d+)?\s*\)", "", name)
    name = re.sub(r"\s+", " ", name).strip(" ()-")
    return name if name else raw.strip()


def _build_score_columns(headers: list, scale_row, id_col, name_col, category: str) -> list:
    score_cols = []
    excluded_patterns = [
        r"^s[\.\s]?no\.?$", r"^sr[\.\s]?no\.?$", r"^sl[\.\s]?no\.?$",
        r"^serial", r"^name$", r"student name", r"full name", r"email", r"phone", r"mobile",
        r"address", r"district", r"qualification", r"college", r"gender", r"dob", r"^age$",
        r"father", r"mother", r"parent", r"reference", r"occupation",
        r"enrolment", r"enrollment", r"^total$",
    ]

    for i, h in enumerate(headers):
        if h is None or str(h).strip() == "":
            continue
        nh = _norm(str(h))

        if any(re.search(p, nh) for p in excluded_patterns):
            continue
        if id_col and _norm(h) == _norm(id_col):
            continue
        if name_col and _norm(h) == _norm(name_col):
            continue
        if category == "attendance":
            continue

        max_val = _extract_max_from_text(h)
        if max_val is None and scale_row is not None and i < len(scale_row) and scale_row[i] is not None:
            max_val = _extract_max_from_text(str(scale_row[i]))
            if max_val is None and _is_numeric(scale_row[i]):
                max_val = float(str(scale_row[i]).replace(",", ""))

        if category == "skill_matrix" and max_val is None:
            nh2 = _norm(str(h))
            if "wpm" in nh2 or "typing" in nh2:
                max_val = None  # raw WPM
            else:
                max_val = 10.0

        score_cols.append({
            "col": str(h).strip(),
            "max": max_val,
            "display_name": _clean_display_name(str(h)),
        })

    return score_cols


# ---------------------------------------------------------------------------
# LLM fallback
# ---------------------------------------------------------------------------

_LLM_CACHE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".schema_llm_cache.json")


def _load_llm_cache() -> dict:
    if os.path.exists(_LLM_CACHE_FILE):
        try:
            with open(_LLM_CACHE_FILE) as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def _save_llm_cache(cache: dict):
    try:
        with open(_LLM_CACHE_FILE, "w") as f:
            json.dump(cache, f, indent=2)
    except Exception as e:
        log.warning(f"Could not save LLM cache: {e}")


def _file_hash(raw_bytes: bytes) -> str:
    return hashlib.sha256(raw_bytes).hexdigest()[:16]


def _ask_llm_for_schema(raw_rows: list, sheet_name: str, real_headers: list):
    try:
        from app.config import settings
        api_key = settings.gemini_api_key
        model = getattr(settings, "gemini_model", "gemini-1.5-flash")
    except Exception:
        api_key = os.environ.get("GEMINI_API_KEY", "")
        model = os.environ.get("GEMINI_MODEL", "gemini-1.5-flash")

    if not api_key or not api_key.startswith("AIza"):
        return None

    rows_text = "\n".join(
        [" | ".join(str(c) if c is not None else "" for c in row) for row in raw_rows[:15]]
    )
    prompt = (
        f"Analyze Excel sheet '{sheet_name}'. First 15 rows:\n{rows_text}\n\n"
        f"Detected headers: {real_headers}\n\n"
        'Reply ONLY with valid JSON: {"header_row":<int>,"id_column":"<name or null>",'
        '"category":"<profile|attendance|assessment|skill_matrix|other_skills|placement|hostelite|parent_meet|irrelevant>",'
        '"score_columns":[{"col":"<exact>","max":<number or null>}],"confidence":<0-1>}\n'
        "col names in score_columns MUST exactly match the detected headers list. "
        "Use irrelevant for uniform/hostel/contact-only sheets. "
        "Never include phone numbers, IDs, or demographics in score_columns."
    )

    try:
        import httpx
        url = (
            f"https://generativelanguage.googleapis.com/v1beta/models/"
            f"{model}:generateContent?key={api_key}"
        )
        payload = {"contents": [{"parts": [{"text": prompt}]}]}
        resp = httpx.post(url, json=payload, timeout=15)
        resp.raise_for_status()
        text = resp.json()["candidates"][0]["content"]["parts"][0]["text"]
        text = text.strip().strip("`").replace("json\n", "", 1).strip()
        result = json.loads(text)
        valid_headers_norm = {_norm(h) for h in real_headers}
        if "score_columns" in result:
            result["score_columns"] = [
                sc for sc in result["score_columns"]
                if _norm(sc.get("col", "")) in valid_headers_norm
            ]
        result["_from_llm"] = True
        return result
    except Exception as e:
        log.warning(f"LLM schema call failed for sheet '{sheet_name}': {e}")
        return None


# ---------------------------------------------------------------------------
# Main detection function
# ---------------------------------------------------------------------------

def detect_schema(source, filename: str = "", llm_confidence_threshold: float = 0.55) -> dict:
    """
    Analyse an Excel file and return a schema manifest dict.
    source: file path (str) or raw bytes.
    """
    if isinstance(source, bytes):
        import io
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True)
        raw_bytes = source
    else:
        wb = openpyxl.load_workbook(source, data_only=True)
        with open(source, "rb") as f:
            raw_bytes = f.read()

    file_hash = _file_hash(raw_bytes)
    llm_cache = _load_llm_cache()
    manifest = {"file": filename or str(source), "file_hash": file_hash, "sheets": {}}

    for sheet_name in wb.sheetnames:
        ws: Any = wb[sheet_name]
        if not hasattr(ws, "iter_rows"):
            continue
        sheet_key = f"{file_hash}::{sheet_name}"

        raw_rows = []
        for row in ws.iter_rows(max_row=20, values_only=True):
            raw_rows.append(list(row))

        if not raw_rows or all(all(v is None for v in r) for r in raw_rows):
            log.info(f"Sheet '{sheet_name}': empty, skipping.")
            continue

        header_row_idx = _detect_header_row(ws, max_scan=20)
        scale_row_idx = _detect_second_header_row(ws, header_row_idx, max_scan=20)

        raw_headers = raw_rows[header_row_idx] if header_row_idx < len(raw_rows) else raw_rows[0]
        scale_row_vals = raw_rows[scale_row_idx] if (scale_row_idx and scale_row_idx < len(raw_rows)) else None

        # Build headers with forward-fill for merged cells, append scale info
        headers = []
        last_filled = ""
        for i, h in enumerate(raw_headers):
            val = str(h).strip() if h is not None else ""
            if val == "" or val.lower() in ("nan", "none"):
                val = last_filled
            else:
                last_filled = val

            if scale_row_vals and i < len(scale_row_vals):
                sv = scale_row_vals[i]
                sv_str = str(sv).strip() if sv is not None else ""
                if sv_str and sv_str.lower() not in ("nan", "none", ""):
                    if re.search(r"out\s+of|^\d+$", sv_str, re.IGNORECASE):
                        if sv_str.lower() not in _norm(val):
                            val = f"{val} ({sv_str})" if val else sv_str
            headers.append(val)

        data_start = (scale_row_idx + 1) if scale_row_idx else (header_row_idx + 1)
        sample_values = raw_rows[data_start: data_start + 5]

        name_col = _find_name_column(headers)
        id_col = _find_id_column(headers)
        category, confidence = _classify_sheet(headers, sample_values, sheet_name)

        from_llm = False
        if confidence < llm_confidence_threshold:
            llm_result = None
            if sheet_key in llm_cache:
                llm_result = llm_cache[sheet_key]
                from_llm = True
            else:
                llm_result = _ask_llm_for_schema(raw_rows, sheet_name, headers)
                if llm_result:
                    llm_cache[sheet_key] = llm_result
                    _save_llm_cache(llm_cache)
                    from_llm = True

            if llm_result:
                category = llm_result.get("category", category)
                confidence = llm_result.get("confidence", confidence)
                if llm_result.get("id_column"):
                    id_col = llm_result["id_column"]
                if "header_row" in llm_result:
                    header_row_idx = llm_result["header_row"]

        total_max_detected = None
        if category in ("irrelevant", "hostelite", "parent_meet", "profile"):
            score_cols = []
        elif category == "attendance":
            score_cols = []
        else:
            plain_headers = [str(h).strip() if h is not None else "" for h in raw_rows[header_row_idx]]
            plain_scale = raw_rows[scale_row_idx] if (scale_row_idx and scale_row_idx < len(raw_rows)) else None
            score_cols = _build_score_columns(plain_headers, plain_scale, id_col, name_col, category)
            if category == "assessment":
                for i, h in enumerate(plain_headers):
                    if h and re.search(r"^total\b", _norm(str(h))):
                        m = re.search(r"(\d+(?:\.\d+)?)", str(h))
                        if m:
                            total_max_detected = float(m.group(1))
                        elif plain_scale and i < len(plain_scale) and plain_scale[i] is not None:
                            m = re.search(r"(\d+(?:\.\d+)?)", str(plain_scale[i]))
                            if m:
                                total_max_detected = float(m.group(1))
                        if total_max_detected is None:
                            total_max_detected = 265.0
                        break

        manifest["sheets"][sheet_name] = {
            "category": category,
            "confidence": round(confidence, 3),
            "header_row": header_row_idx,
            "scale_row": scale_row_idx,
            "id_column": id_col,
            "name_column": name_col,
            "score_columns": score_cols,
            "total_max": total_max_detected,
            "from_llm": from_llm,
        }
        log.info(
            f"Sheet '{sheet_name}': category={category} conf={confidence:.2f} "
            f"hdr={header_row_idx} scale={scale_row_idx} name={name_col} id={id_col} "
            f"scores={len(score_cols)} llm={from_llm}"
        )

    return manifest


def print_manifest(manifest: dict):
    print(f"\n{'='*70}")
    print(f"SCHEMA MANIFEST - {manifest['file']}  (hash: {manifest.get('file_hash','?')})")
    print(f"{'='*70}")
    for sheet, info in manifest["sheets"].items():
        print(f"\n  Sheet: '{sheet}'")
        print(f"    Category    : {info['category']}  (conf={info['confidence']:.2f}{'  [LLM]' if info.get('from_llm') else ''})")
        print(f"    Header row  : {info['header_row']}  |  Scale row: {info.get('scale_row')}")
        print(f"    Name col    : {info.get('name_column')}")
        print(f"    ID col      : {info.get('id_column')}")
        if info.get("score_columns"):
            print(f"    Score cols ({len(info['score_columns'])}):")
            for sc in info["score_columns"]:
                mx = f"max={sc['max']}" if sc.get("max") is not None else "max=?"
                print(f"      - {sc['col']!r:48s} {mx:12s} display={sc.get('display_name','')!r}")
        else:
            print(f"    Score cols  : (none)")
    print(f"\n{'='*70}\n")


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    if len(sys.argv) < 2:
        default = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "Students Complete Details - CCDP 2(1).xlsx"
        )
        path = default if os.path.exists(default) else None
        if not path:
            print("Usage: python schema_detector.py <path_to_excel.xlsx>")
            sys.exit(1)
    else:
        path = sys.argv[1]

    print(f"Analysing: {path}")
    manifest = detect_schema(path, os.path.basename(path))
    print_manifest(manifest)

    out_path = os.path.splitext(path)[0] + "_schema_manifest.json"
    with open(out_path, "w") as fh:
        json.dump(manifest, fh, indent=2)
    print(f"Manifest saved to: {out_path}")
